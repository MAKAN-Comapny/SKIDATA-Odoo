"""
Dry-run the SharePoint -> Odoo import hierarchy WITHOUT creating invoices.

Flow:
  1. Authenticate to SharePoint (Graph)
  2. Open folder SKIDATA - Odoo / SKIDATA-Odoo
  3. Download the latest *Parking-Transactions-Odoo* file
  4. Read rows and preview the payloads that WOULD be sent to Odoo
  5. Stop. No Odoo XML-RPC create/update.
"""
import json
import os
import sys
from datetime import date, datetime

from dotenv import load_dotenv

import sharepoint_download

PREVIEW_ROWS = 5


def _load_env():
    here = os.path.dirname(os.path.abspath(__file__))
    parent = os.path.dirname(here)
    load_dotenv(os.path.join(parent, ".env"))
    load_dotenv(os.path.join(here, ".env"), override=True)


def _fmt(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return ""
    return str(value)


def load_rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        return [
            {h: _fmt(v) for h, v in zip(headers, row)}
            for row in ws.iter_rows(min_row=2, values_only=True)
        ]

    import csv

    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def preview_payload(row, odoo_cfg):
    collected = float(row.get("Invoice Lines/Unit Price") or 0)
    analytic_raw = row.get("Invoice Lines/Analytics Distribution") or "{}"
    try:
        analytic = json.loads(analytic_raw)
    except json.JSONDecodeError:
        analytic = {"_parse_error": analytic_raw}

    bill_date = (row.get("Invoice/Bill Date") or "")[:10]
    quantity = row.get("Invoice Lines/Quantity", row.get("Quantity", 1))

    return {
        "move_type": "out_invoice",
        "partner_id": odoo_cfg["walkin_partner_id"],
        "journal_id": odoo_cfg["journal_id"],
        "invoice_date": bill_date,
        "in_time": row.get("In Time"),
        "out_time": row.get("Out Time"),
        "ticket_sequence": row.get("Ticket Sequence"),
        "payment_reference": row.get("Payment Reference"),
        "plate_number": row.get("Plate Number"),
        "invoice_line": {
            "account_id": odoo_cfg["sales_account_id"],
            "name": row.get("Invoice Lines/Label"),
            "quantity": quantity,
            "price_unit": collected,
            "tax_ids": [odoo_cfg["tax_id"]] if odoo_cfg["tax_id"] else [],
            "analytic_distribution": analytic,
        },
    }


def odoo_preview_config():
    def optional_int(name):
        raw = os.getenv(name)
        if not raw or not str(raw).strip():
            return None
        return int(str(raw).strip().strip('"'))

    return {
        "url": os.getenv("ODOO_URL"),
        "db": os.getenv("ODOO_DB"),
        "username": os.getenv("ODOO_USERNAME"),
        "journal_id": optional_int("ODOO_JOURNAL_ID"),
        "sales_account_id": optional_int("ODOO_SALES_ACCOUNT_ID"),
        "tax_id": optional_int("ODOO_TAX_ID"),
        "walkin_partner_id": optional_int("ODOO_WALKIN_PARTNER_ID"),
    }


def main(preview_limit=PREVIEW_ROWS):
    _load_env()

    print("=" * 64)
    print("DRY RUN: SharePoint latest file -> Odoo payload preview")
    print("No invoices will be created.")
    print("=" * 64)

    print("\n[1/4] Connecting to SharePoint...")
    local_path, item, folder = sharepoint_download.fetch_latest_export(dest_dir="downloads")
    if local_path is None:
        print("FAILED: no matching export found.")
        print("Tried folders:", folder)
        sys.exit(1)

    print(f"  Folder used : {folder}")
    print(f"  Latest file : {item.get('name')}")
    print(f"  Modified    : {item.get('lastModifiedDateTime')}")
    print(f"  Downloaded  : {os.path.abspath(local_path)}")
    print(f"  Size bytes  : {os.path.getsize(local_path)}")

    print("\n[2/4] Reading mapped Odoo file (no extra transform)...")
    rows = load_rows(local_path)
    print(f"  Rows        : {len(rows)}")
    if rows:
        print("  Columns     :")
        for col in rows[0].keys():
            print(f"    - {col}")

    print("\n[3/4] Building preview payloads (would be sent to Odoo)...")
    cfg = odoo_preview_config()
    print(
        "  Odoo target :"
        f" db={cfg['db'] or '(not set)'} url={cfg['url'] or '(not set)'}"
    )
    print(
        "  IDs         :"
        f" journal={cfg['journal_id']} account={cfg['sales_account_id']}"
        f" tax={cfg['tax_id']} walkin={cfg['walkin_partner_id']}"
    )
    print("  Price rule  : use collected amount as-is (10, 30, ...)")
    print("  Also sent   : Payment Reference, Plate Number")

    sample = rows[:preview_limit]
    for idx, row in enumerate(sample, start=1):
        payload = preview_payload(row, cfg)
        line = payload["invoice_line"]
        print(f"\n  Preview {idx}/{len(sample)}")
        print(f"    ticket_sequence : {payload['ticket_sequence']}")
        print(f"    invoice_date    : {payload['invoice_date']}")
        print(f"    label           : {line['name']}")
        print(f"    price_unit      : {line['price_unit']}")
        print(f"    payment_ref     : {payload['payment_reference']}")
        print(f"    plate_number    : {payload['plate_number']}")
        print(f"    analytics       : {line['analytic_distribution']}")
        print(f"    location/region : {row.get('Location')} / {row.get('Region')}")
        print(f"    service/sub     : {row.get('Service')} / {row.get('Sub Service')}")

    print("\n[4/4] Odoo write step SKIPPED (dry-run).")
    print("Hierarchy OK: SharePoint latest file -> read mapped rows -> preview payloads.")
    print("=" * 64)


if __name__ == "__main__":
    limit = int(sys.argv[1]) if len(sys.argv) > 1 else PREVIEW_ROWS
    main(limit)
