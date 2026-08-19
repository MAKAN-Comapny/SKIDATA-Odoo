"""
Import a batch of tickets from the pre-mapped Odoo-import-style file
into Odoo as customer invoices.

Unit Price in the source file is the collected amount (10, 30, ...).
It is sent to Odoo as-is (not VAT-excluded / net).

Usage:
  python import_csv_batch.py path/to/file.xlsx [limit]
  python import_csv_batch.py path/to/file.xlsx 3 --yes
  python import_csv_batch.py path/to/file.xlsx 3 --post
  python import_csv_batch.py path/to/file.xlsx 3 --post --pay --yes
"""
import argparse
import csv
import json
from datetime import date, datetime

import config
import odoo_client


def _clean_text(value):
    text = _fmt(value).strip()
    if text.lower() in ("", "nan", "none", "-", "#"):
        return False
    return text


def _fmt(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ""
    return str(v)


def load_rows(path, limit=None):
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        rows = [
            {h: _fmt(v) for h, v in zip(headers, row)}
            for row in ws.iter_rows(min_row=2, values_only=True)
        ]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
    return rows[:limit] if limit else rows


def build_payload(row):
    price_unit = float(row["Invoice Lines/Unit Price"])

    try:
        analytic = json.loads(row["Invoice Lines/Analytics Distribution"] or "{}")
    except json.JSONDecodeError:
        analytic = {}
    analytic = {str(k): v for k, v in analytic.items()}

    invoice_date = _fmt(row["Invoice/Bill Date"])[:10]
    ticket_sequence = str(row["Ticket Sequence"]).strip()

    payload = {
        "move_type": "out_invoice",
        "partner_id": config.ODOO_WALKIN_PARTNER_ID,
        "journal_id": config.ODOO_JOURNAL_ID,
        "invoice_date": invoice_date,
        "in_time": row["In Time"],
        "out_time": row["Out Time"],
        "ticket_sequence": ticket_sequence,
        "payment_reference": _clean_text(row.get("Payment Reference")),
        "plate_number": _clean_text(row.get("Plate Number")),
        "invoice_line_ids": [
            (
                0,
                0,
                {
                    "account_id": config.ODOO_SALES_ACCOUNT_ID,
                    "name": row["Invoice Lines/Label"],
                    "quantity": float(
                        row.get("Invoice Lines/Quantity", row.get("Quantity", 1))
                    ),
                    "price_unit": price_unit,
                    "tax_ids": [(6, 0, [config.ODOO_TAX_ID])],
                    "analytic_distribution": analytic,
                },
            )
        ],
    }
    return payload


def parse_args():
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("csv_path")
    parser.add_argument("limit", nargs="?", type=int, default=10)
    parser.add_argument("--yes", action="store_true", help="Skip confirmation prompt.")
    parser.add_argument(
        "--post", action="store_true", help="Validate (post) each created invoice."
    )
    parser.add_argument(
        "--pay",
        action="store_true",
        help="Register full payment for each invoice (implies --post).",
    )
    return parser.parse_args()


def main(csv_path, limit=10, auto_yes=False, post=False, pay=False):
    payment_journal_id = None
    if pay:
        post = True
        payment_journal_id = config.ODOO_PAYMENT_JOURNAL_ID
        if not payment_journal_id:
            raise RuntimeError(
                "--pay requires ODOO_PAYMENT_JOURNAL_ID to be set "
                "(the bank/cash journal these tickets settle through)."
            )

    rows = load_rows(csv_path, limit)
    action_desc = "POSTED + PAID" if pay else "POSTED" if post else "DRAFT"
    print(f"Loaded {len(rows)} row(s) from {csv_path} (limit={limit}).")
    print(f"Odoo: {config.ODOO_URL} / {config.ODOO_DB}")
    print(f"Mode: {action_desc}\n")

    for r in rows:
        print(
            f"  - ticket {r['Ticket Sequence']} | {r['Invoice Lines/Label']} "
            f"| amount {r['Invoice Lines/Unit Price']} "
            f"| ref {r.get('Payment Reference') or '-'} "
            f"| plate {r.get('Plate Number') or '-'}"
        )

    if not auto_yes:
        confirm = input(f"\nCreate {len(rows)} {action_desc} invoice(s) in Odoo now? [y/N] ")
        if confirm.strip().lower() != "y":
            print("Aborted, nothing created.")
            return
    else:
        print(f"\nAuto-confirm: creating {len(rows)} {action_desc} invoice(s)...")

    ticket_payloads = [(str(r["Ticket Sequence"]), build_payload(r)) for r in rows]

    client = odoo_client.OdooClient()
    results = client.process_batch(
        ticket_payloads, post=post, pay=pay, payment_journal_id=payment_journal_id
    )

    print("\n=== Results ===")
    print("Created:", results["created"])
    print("Skipped (duplicate):", results["skipped_duplicate"])
    print("Failed:", results["failed"])


if __name__ == "__main__":
    args = parse_args()
    main(args.csv_path, args.limit, auto_yes=args.yes, post=args.post, pay=args.pay)
