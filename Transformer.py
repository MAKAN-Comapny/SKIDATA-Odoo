"""
Transform Parking Transactions export into Odoo-ready invoice lines.

Steps:
  1. Load Parking-Transactions .xlsx from exports/
  2. Keep rows where Payment is not 0 / blank
  3. Map/rename columns to the Odoo structure
  4. Write result to Cleaned-Exports/ and clear exports/
"""
import os
import glob
import re
from decimal import Decimal, InvalidOperation

import pandas as pd
from openpyxl import Workbook

EXPORTS_DIR = "exports"
CLEANED_EXPORTS_DIR = "Cleaned-Exports"
PARKING_FILE_MARKER = "Parking-Transactions"

# Exact output column names (Odoo import structure)
OUT_PARTNER = "Partner"
OUT_INVOICE_BILL_DATE = "Invoice/Bill Date"
OUT_INVOICE_LABEL = "Invoice Lines/Label"
OUT_QUANTITY = "Quantity"
OUT_UNIT_PRICE = "Invoice Lines/Unit Price"
OUT_TAXES = "invoice Lines/Taxes"
OUT_IN_TIME = "In Time"
OUT_OUT_TIME = "Out Time"
OUT_TICKET_SEQUENCE = "Ticket Sequence"
OUT_ANALYTICS = "Invoice Lines/Analytics Distribution"
OUT_PLATE_NUMBER = "Plate Number"
OUT_LOCATION = "Location"
OUT_REGION = "Region"
OUT_SERVICE = "Service"
OUT_SUB_SERVICE = "Sub Service"

FIXED_PARTNER = "SKIDATA"
FIXED_QUANTITY = 1
FIXED_TAXES = "Sales Tax"
FIXED_ANALYTICS = '{"paid parking":100}'

# SKIDATA Car Park name -> Odoo Location / Region / Service / Sub Service
CAR_PARK_MAPPING = {
    "Westfield DAMMAM": {
        "location": "02 - Nakheel Mall Dammam",
        "region": "Dammam Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "View Mall": {
        "location": "01 - The View Mall Riyadh",
        "region": "Riyadh Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "Alnoor Mall": {
        "location": "06 - Al Noor Mall Madinah",
        "region": "Madinah Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "Mall of Arabia": {
        "location": "03 - Arab Mall Jeddah",
        "region": "Jeddah Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "Premium Carpark": {
        "location": "06 - Al Noor Mall Madinah",
        "region": "Madinah Region",
        "service": "Paid Parking",
        "sub_service": "Premium",
    },
}

OUTPUT_COLUMNS = [
    OUT_PARTNER,
    OUT_INVOICE_BILL_DATE,
    OUT_INVOICE_LABEL,
    OUT_QUANTITY,
    OUT_UNIT_PRICE,
    OUT_TAXES,
    OUT_IN_TIME,
    OUT_OUT_TIME,
    OUT_TICKET_SEQUENCE,
    OUT_ANALYTICS,
    OUT_PLATE_NUMBER,
    OUT_LOCATION,
    OUT_REGION,
    OUT_SERVICE,
    OUT_SUB_SERVICE,
]


def _normalize_header(name: str) -> str:
    """Lowercase and collapse separators for fuzzy column matching."""
    return re.sub(r"[\s/_-]+", "", str(name).strip().lower())


def _find_column(columns, *candidates: str) -> str:
    """Return the actual column name matching any candidate."""
    normalized = {_normalize_header(c): c for c in columns}
    for candidate in candidates:
        key = _normalize_header(candidate)
        if key in normalized:
            return normalized[key]
    raise KeyError(
        f"Could not find column matching any of {candidates}. "
        f"Available columns: {list(columns)}"
    )


def _is_blank(value) -> bool:
    """Treat empty / NaN / SKIDATA plate placeholders (-, #) as missing."""
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    text = str(value).strip()
    return text == "" or text.lower() == "nan" or text in ("-", "#")


def _pick_plate(entry_plate, exit_plate) -> str:
    """Use whichever plate exists; if both exist take either; else blank."""
    entry_blank = _is_blank(entry_plate)
    exit_blank = _is_blank(exit_plate)
    if not entry_blank:
        return str(entry_plate).strip()
    if not exit_blank:
        return str(exit_plate).strip()
    return ""


def _serial_as_text(value) -> str:
    """Keep full serial/ticket sequence as text (never as Excel number)."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        as_int = int(round(value))
        if abs(value - as_int) < 1e-9:
            return str(as_int)
        return format(value, "f").rstrip("0").rstrip(".")
    text = str(value).strip()
    if text == "" or text.lower() == "nan":
        return ""
    if re.fullmatch(r"[+-]?\d+(\.\d+)?[eE][+-]?\d+", text):
        try:
            d = Decimal(text)
            if d == d.to_integral_value():
                return format(d.to_integral_value(), "f")
            return format(d, "f")
        except InvalidOperation:
            return text
    if re.fullmatch(r"[+-]?\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _payment_is_nonzero(series: pd.Series) -> pd.Series:
    """True where Payment has a numeric value other than 0."""
    numeric = pd.to_numeric(series, errors="coerce")
    return numeric.notna() & (numeric != 0)


def _map_car_park(car_park_value) -> dict:
    """Map SKIDATA Car Park to Odoo Location / Region / Service / Sub Service."""
    key = "" if car_park_value is None or (isinstance(car_park_value, float) and pd.isna(car_park_value)) else str(car_park_value).strip()
    mapped = CAR_PARK_MAPPING.get(key)
    if mapped is None:
        return {
            "location": "",
            "region": "",
            "service": "",
            "sub_service": "",
        }
    return mapped


def _write_odoo_excel(df: pd.DataFrame, dest_path: str) -> None:
    """Write workbook with Ticket Sequence forced as Excel text."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = list(df.columns)
    ws.append(headers)
    ticket_col_idx = headers.index(OUT_TICKET_SEQUENCE) + 1  # 1-based

    for row in df.itertuples(index=False, name=None):
        values = []
        for col_name, value in zip(headers, row):
            if col_name == OUT_TICKET_SEQUENCE:
                values.append(_serial_as_text(value))
            elif pd.isna(value):
                values.append(None)
            else:
                values.append(value)
        ws.append(values)

    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=ticket_col_idx)
        if cell.value is None:
            continue
        cell.value = str(cell.value)
        cell.number_format = "@"

    os.makedirs(os.path.dirname(dest_path) or ".", exist_ok=True)
    wb.save(dest_path)


def transform_parking_file(src_path: str, dest_path: str) -> tuple[int, int]:
    """
    Filter + remap one Parking Transactions workbook.
    Returns (rows_before_filter, rows_after_filter).
    """
    # Resolve Serial Number column, then re-read it as text to preserve digits.
    header_df = pd.read_excel(src_path, nrows=0)
    serial_col = _find_column(header_df.columns, "Serial Number", "SerialNumber")
    df = pd.read_excel(src_path, dtype={serial_col: str})
    rows_before = len(df)

    payment_col = _find_column(df.columns, "Payment")
    txn_date_col = _find_column(
        df.columns,
        "Transaction Date/Time",
        "Transaction Date / Time",
        "Transaction Date Time",
    )
    article_col = _find_column(df.columns, "Article")
    entry_date_col = _find_column(
        df.columns, "Entry Date / Time", "Entry Date/Time", "Entry Date Time"
    )
    exit_date_col = _find_column(
        df.columns, "Exit Date / Time", "Exit Date/Time", "Exit Date Time"
    )
    entry_plate_col = _find_column(
        df.columns, "Entry Plate Number", "Entry Plate", "EntryPlateNumber"
    )
    exit_plate_col = _find_column(
        df.columns, "Exit Plate Number", "Exit Plate", "ExitPlateNumber"
    )
    car_park_col = _find_column(df.columns, "Car Park", "CarPark")

    filtered = df[_payment_is_nonzero(df[payment_col])].copy()
    rows_after = len(filtered)

    location_maps = [_map_car_park(v) for v in filtered[car_park_col].values]

    out = pd.DataFrame(
        {
            OUT_PARTNER: FIXED_PARTNER,
            OUT_INVOICE_BILL_DATE: filtered[txn_date_col].values,
            OUT_INVOICE_LABEL: filtered[article_col].values,
            OUT_QUANTITY: FIXED_QUANTITY,
            OUT_UNIT_PRICE: pd.to_numeric(filtered[payment_col], errors="coerce").values,
            OUT_TAXES: FIXED_TAXES,
            OUT_IN_TIME: filtered[entry_date_col].values,
            OUT_OUT_TIME: filtered[exit_date_col].values,
            OUT_TICKET_SEQUENCE: [
                _serial_as_text(v) for v in filtered[serial_col].values
            ],
            OUT_ANALYTICS: FIXED_ANALYTICS,
            OUT_PLATE_NUMBER: [
                _pick_plate(e, x)
                for e, x in zip(
                    filtered[entry_plate_col].values,
                    filtered[exit_plate_col].values,
                )
            ],
            OUT_LOCATION: [m["location"] for m in location_maps],
            OUT_REGION: [m["region"] for m in location_maps],
            OUT_SERVICE: [m["service"] for m in location_maps],
            OUT_SUB_SERVICE: [m["sub_service"] for m in location_maps],
        }
    )
    out = out[OUTPUT_COLUMNS]
    _write_odoo_excel(out, dest_path)
    return rows_before, rows_after


def run():
    """Transform Parking Transactions files into Odoo format under Cleaned-Exports/."""
    all_files = glob.glob(os.path.join(EXPORTS_DIR, "*.xlsx"))
    parking_files = [f for f in all_files if PARKING_FILE_MARKER in os.path.basename(f)]

    if not parking_files:
        raise SystemExit(
            f"No Parking Transactions .xlsx files found in '{EXPORTS_DIR}'. "
            "Run scraper first."
        )

    os.makedirs(CLEANED_EXPORTS_DIR, exist_ok=True)

    for src in parking_files:
        base = os.path.basename(src)
        if base.lower().endswith(".xlsx"):
            dest_name = base[:-5] + "-Odoo.xlsx"
        else:
            dest_name = base + "-Odoo.xlsx"
        dest = os.path.join(CLEANED_EXPORTS_DIR, dest_name)

        before, after = transform_parking_file(src, dest)
        print(f"  {base}")
        print(f"    rows: {before} -> {after} (Payment != 0)")
        print(f"    wrote: {dest_name}")

    for file in all_files:
        os.remove(file)

    print(
        f"Transformed {len(parking_files)} Parking Transactions file(s) "
        f"to {CLEANED_EXPORTS_DIR}/"
    )


if __name__ == "__main__":
    run()
